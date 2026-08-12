"""§7 PPT 출력 — 슬라이드 크기·페이지 순서·표 페이지·병합 순서.

확정 사양 중 이 파일이 지키는 것:
  §7.1 슬라이드는 항상 16:9. 표 때문에 덱을 넓히면 plot까지 늘어진다
  §7.2 plot 전부 → 표 전부(한 벌) → 제외 이력. 실험마다 표를 붙이면 중복이다
  §7.3 표는 전용 페이지 · CAT1마다 한 장 · 9pt 유지(overflow) 또는 12장 분할
  §7.4 회색 헤더·교대 행·규격 이탈 셀만 붉게
  §10.5 **병합 먼저, 값 나중** — 반대로 하면 "PA1\\nPA1\\nPA1"이 된다
"""
from __future__ import annotations

import polars as pl
import pytest
from pptx.oxml.ns import qn

from etreport.data.reformatter import Reformatter, Rule
from etreport.model.specs import PageSpec, PlotSpec, ReportSpec
from etreport.render import pptgen
from etreport.render.pptgen import TableData

RATIO_16_9 = 16 / 9


def _rf() -> Reformatter:
    return Reformatter(rules=[
        Rule(category="REAL", itemid="P1", alias="Vt", absolute=False, scale=1.0,
             formula="", unit="V", speclow=0.3, spechigh=0.6, target=0.45, row=2),
        Rule(category="REAL", itemid="P2", alias="Ioff", absolute=True, scale=1.0,
             formula="", unit="A", speclow=None, spechigh=None, target=None, row=3),
    ])


def _report(n_pages: int = 1) -> ReportSpec:
    pages = []
    for i in range(1, n_pages + 1):
        page = PageSpec(number=i, title=f"P{i}")
        page.slots[0] = PlotSpec(title="Vt-Ioff", x="Vt", y="Ioff")
        pages.append(page)
    return ReportSpec(report="R", pages=pages)


def _table(name: str = "DC", n_wafers: int = 4, n_rows: int = 3) -> TableData:
    wafers = [f"{i:02d}" for i in range(1, n_wafers + 1)]
    half = max(1, len(wafers) // 2)
    header = [("PA1", wafers[:half]), ("PB2", wafers[half:])]
    header = [(lot, ws) for lot, ws in header if ws]
    rows = [{"cats": ["NMOS" if i < 2 else "PMOS", "SVT"],
             "item": f"item{i}", "values": [0.4 + i] * len(wafers),
             "offspec": [i == 1] + [False] * (len(wafers) - 1)}
            for i in range(n_rows)]
    return TableData(name, header, rows)


def _deck(tables: list[TableData], experiments=("",), mode="overflow",
          n_pages: int = 1, report: ReportSpec | None = None):
    return pptgen.build_deck(
        report=report or _report(n_pages),
        experiments=list(experiments),
        group_styles_of=lambda exp: [],
        plot_data_of=lambda exp, spec: {},
        tables=tables,
        rf=_rf(),
        log_patterns=["Ioff*"],
        exclusion_log=pl.DataFrame({"key_hash": [], "reason": [], "created_at": []}),
        table_mode=mode,
    )


def _kinds(prs) -> list[str]:
    """슬라이드 종류 목록 — plot / table / excl."""
    out = []
    for s in prs.slides:
        title = next((sh.text_frame.text for sh in s.shapes
                      if sh.has_text_frame and sh.text_frame.text), "")
        if any(sh.has_table for sh in s.shapes):
            out.append("table")
        elif title.startswith("제외 포인트 이력"):
            out.append("excl")
        else:
            out.append("plot")
    return out


def _first_table(prs, idx: int):
    slide = prs.slides[idx]
    return next(sh.table for sh in slide.shapes if sh.has_table)


def _title(prs, idx: int) -> str:
    return next(sh.text_frame.text for sh in prs.slides[idx].shapes
                if sh.has_text_frame and sh.text_frame.text)


# ── §7.1 슬라이드 크기 ───────────────────────────────────────
@pytest.mark.parametrize("n_wafers,mode", [(4, "overflow"), (25, "overflow"),
                                           (25, "split"), (60, "overflow")])
def test_slide_is_always_16_9(n_wafers, mode):
    """★ wafer가 몇 장이든 비율은 1.778 — 덱을 넓히면 plot이 가로로 늘어진다."""
    prs = _deck([_table(n_wafers=n_wafers)], mode=mode)

    assert prs.slide_width.inches == pytest.approx(13.333)
    assert prs.slide_height.inches == pytest.approx(7.5)
    assert prs.slide_width / prs.slide_height == pytest.approx(RATIO_16_9, rel=1e-3)


def test_legacy_wide_mode_is_treated_as_overflow():
    """예전 설정 파일의 'wide'는 overflow와 같다 — 덱을 넓히지 않는다(§10.8)."""
    assert pptgen.table_mode_of("wide") == "overflow"
    prs = _deck([_table(n_wafers=25)], mode="wide")
    assert prs.slide_width.inches == pytest.approx(13.333)


# ── §7.2 페이지 순서·중복 ────────────────────────────────────
def test_page_order_is_plots_then_tables_then_exclusions():
    """★ plot 전부 → 표 전부 → 이력."""
    prs = _deck([_table("DC"), _table("AC")], experiments=["실험A"], n_pages=2)

    assert _kinds(prs) == ["plot", "plot", "table", "table", "excl"]


def test_tables_are_not_repeated_per_experiment():
    """★ 표는 (lot,wafer)별 집계라 실험과 무관 — 실험이 셋이어도 한 벌만."""
    prs = _deck([_table("DC"), _table("AC")],
                experiments=["A", "B", "C"], n_pages=2)

    kinds = _kinds(prs)
    assert kinds.count("plot") == 6            # 실험 3 × 페이지 2
    assert kinds.count("table") == 2           # CAT1 2개 — 실험 수와 무관
    assert kinds[-1] == "excl"


def test_legacy_table_slot_does_not_duplicate_the_table():
    """예전 템플릿의 Type=table 행이 있어도 표는 전용 페이지에만 나온다(§7.3)."""
    rep = _report()
    rep.pages[0].slots[1] = PlotSpec(title="표", x="", y="", type="table")

    prs = _deck([_table("DC")], report=rep)

    assert _kinds(prs).count("table") == 1


# ── §10.5 병합 순서 ──────────────────────────────────────────
def test_lot_header_is_merged_before_text():
    """★ lot 헤더는 'PA1' 하나여야 한다 ('PA1\\nPA1\\nPA1'이 아니라)."""
    prs = _deck([_table(n_wafers=6)])
    tbl = _first_table(prs, 1)

    assert tbl.cell(0, 3).text == "PA1"
    assert "\n" not in tbl.cell(0, 3).text
    assert tbl.cell(0, 3).span_width == 3          # wafer 3장을 덮는다
    assert tbl.cell(0, 6).text == "PB2"


def test_cat_columns_merge_and_break_on_upper_change():
    """CAT2가 세로 병합되고, 상위가 바뀌면 하위 병합도 끊긴다(§3.3)."""
    td = TableData("DC", [("PA1", ["01"])], [
        {"cats": ["NMOS", "SVT"], "item": "a", "values": [1.0],
         "offspec": [False]},
        {"cats": ["NMOS", "SVT"], "item": "b", "values": [2.0],
         "offspec": [False]},
        {"cats": ["PMOS", "SVT"], "item": "c", "values": [3.0],
         "offspec": [False]},
    ])
    tbl = _first_table(_deck([td]), 1)

    assert tbl.cell(2, 0).text == "NMOS"
    assert tbl.cell(2, 0).span_height == 2         # 두 행 병합
    assert tbl.cell(4, 0).text == "PMOS"
    # CAT3는 값이 같아도 CAT2가 바뀌면 따로 — 3행 전체가 하나로 묶이지 않는다
    assert tbl.cell(2, 1).span_height == 2
    assert tbl.cell(4, 1).text == "SVT"


# ── §7.3 폭·모드 ─────────────────────────────────────────────
def test_overflow_keeps_9pt_and_runs_off_the_slide():
    """wafer가 많으면 글자를 줄이지 않고 표가 슬라이드 밖으로 나간다."""
    total, wafer_w, overflow = pptgen.table_width_in(30)
    assert overflow
    assert wafer_w == pytest.approx(pptgen.WAFER_W_IN)   # 9pt 폭 유지
    assert total > pptgen.BASE_W_IN

    prs = _deck([_table(n_wafers=30)])
    frame = next(sh for sh in prs.slides[1].shapes if sh.has_table)
    assert frame.width > prs.slide_width
    assert pptgen.OVERFLOW_NOTE in _title(prs, 1)       # 제목에 안내 문구
    sizes = {run.font.size for row in frame.table.rows for cell in row.cells
             for para in cell.text_frame.paragraphs for run in para.runs}
    assert sizes == {pptgen.FONT_MIN}


def test_small_table_fits_and_says_nothing():
    prs = _deck([_table(n_wafers=4)])
    frame = next(sh for sh in prs.slides[1].shapes if sh.has_table)

    assert frame.width <= prs.slide_width
    assert pptgen.OVERFLOW_NOTE not in _title(prs, 1)


def test_split_mode_breaks_into_slides_of_12():
    """split은 wafer 12장씩 나누고 제목에 (1/N)을 붙인다."""
    prs = _deck([_table(n_wafers=25)], mode="split")

    kinds = _kinds(prs)
    assert kinds.count("table") == 3                   # 12 + 12 + 1
    assert "(1/3)" in _title(prs, 1) and "(3/3)" in _title(prs, 3)
    frame = next(sh for sh in prs.slides[1].shapes if sh.has_table)
    assert frame.width <= prs.slide_width              # 나눴으니 넘치지 않는다


# ── §7.4 스타일 ──────────────────────────────────────────────
def test_offspec_cells_are_the_only_red_ones():
    prs = _deck([_table(n_wafers=4)])
    tbl = _first_table(prs, 1)

    assert tbl.cell(3, 3).fill.fore_color.rgb == pptgen.SPEC_FILL   # 규격 이탈
    assert tbl.cell(3, 4).fill.fore_color.rgb != pptgen.SPEC_FILL
    assert tbl.cell(2, 3).fill.fore_color.rgb != pptgen.SPEC_FILL
    assert not tbl.first_row and not tbl.horz_banding   # 기본 파란 줄무늬 제거


def test_cell_border_xml_is_in_schema_order():
    """a:lnB는 tcPr의 맨 앞 — 순서가 어긋나면 PowerPoint가 '복구' 창을 띄운다."""
    prs = _deck([_table(n_wafers=4)])
    tbl = _first_table(prs, 1)
    tc_pr = tbl.cell(1, 0)._tc.get_or_add_tcPr()

    tags = [child.tag for child in tc_pr]
    assert tags[0] == qn("a:lnB")
    assert any(t.endswith("}solidFill") for t in tags[1:])          # 채우기는 뒤


def test_deck_round_trips_through_a_file(tmp_path):
    """저장·재오픈이 되는지 — XML을 직접 만졌으므로 한 번은 확인한다."""
    from pptx import Presentation as Open

    p = tmp_path / "deck.pptx"
    _deck([_table("DC", n_wafers=14), _table("AC")], experiments=["A", "B"]).save(p)

    again = Open(str(p))
    assert len(again.slides) == 2 * 1 + 2 + 1
    assert again.slide_width / again.slide_height == pytest.approx(RATIO_16_9,
                                                                  rel=1e-3)
