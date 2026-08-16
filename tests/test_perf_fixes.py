"""성능 수정 회귀 — 느려지는 길로 되돌아가지 않게 고정한다.

여기서 지키는 것 넷:
  1. 표의 열 폭·행 높이는 `pptgen.set_grid`로 준다. python-pptx의
     `columns[i].width` setter는 열 수의 **제곱**으로 느려진다(wafer 344장에서
     표 한 장 57.8초 → 25.0초). 결과가 setter 경로와 같은지 값으로 비교한다.
  2. trend는 wafer 집계를 **그룹당 한 번**만 한다. item마다 부르면 group_by가
     item 수만큼 반복된다(alias 25개 기준 19배).
  3. xlsx 서식은 셀이 아니라 **구간(run)** 단위다 — COM 왕복이 행 수준으로 준다.
  4. 오래 걸리는 작업은 진행을 보고한다(분모가 도중에 바뀌지 않는다).
"""
from __future__ import annotations

import polars as pl
from pptx import Presentation
from pptx.util import Inches

from etreport.data.reformatter import Reformatter, Rule
from etreport.export.excel import flag_runs, format_runs, number_format
from etreport.model.specs import GroupStyle, PageSpec, PlotSpec, ReportSpec
from etreport.render import mpl_renderer, pptgen
from etreport.render.pptgen import TableData


def _rf() -> Reformatter:
    return Reformatter(rules=[
        Rule(category="REAL", itemid="P1", alias="A", absolute=False, scale=1.0,
             formula="", unit="V", speclow=0.3, spechigh=0.6, target=0.45,
             row=2, w=0.34, l=12.5),
        Rule(category="REAL", itemid="P2", alias="B", absolute=False, scale=1.0,
             formula="", unit="V", speclow=None, spechigh=None, target=None,
             row=3, w=0.66, l=25.0),
    ])


def _table(n_wafers: int = 6, n_rows: int = 4) -> TableData:
    ws = [f"{i:02d}" for i in range(1, n_wafers + 1)]
    half = max(1, len(ws) // 2)
    header = [(lot, w) for lot, w in
              (("PA1", ws[:half]), ("PB2", ws[half:])) if w]
    rows = [{"cats": ["NMOS" if i < 2 else "PMOS", "SVT"], "item": f"it{i}",
             "values": [0.4 + i] * n_wafers,
             "offspec": [i == 1] + [False] * (n_wafers - 1)}
            for i in range(n_rows)]
    return TableData("DC", header, rows)


def _prs():
    prs = Presentation()
    prs.slide_width = Inches(pptgen.BASE_W_IN)
    prs.slide_height = Inches(pptgen.BASE_H_IN)
    return prs, prs.slide_layouts[6]


# ── 1) 표 그리드 ─────────────────────────────────────────────
def test_set_grid_matches_pptx_setters():
    """set_grid의 결과 == python-pptx setter를 하나씩 쓴 결과."""
    widths = [Inches(1.15), Inches(1.35), Inches(0.62), Inches(0.62)]
    row_h = Inches(0.24)

    prs, layout = _prs()
    slow = prs.slides.add_slide(layout).shapes.add_table(
        3, 4, Inches(0.35), Inches(1.0), Inches(3.74), Inches(0.72))
    for i, w in enumerate(widths):                 # 예전 경로(느린 쪽)
        slow.table.columns[i].width = w
    for r in range(3):
        slow.table.rows[r].height = row_h

    fast = prs.slides.add_slide(layout).shapes.add_table(
        3, 4, Inches(0.35), Inches(1.0), Inches(3.74), Inches(0.72))
    pptgen.set_grid(fast, widths, row_h)

    assert [c.width for c in fast.table.columns] == \
           [c.width for c in slow.table.columns]
    assert [r.height for r in fast.table.rows] == \
           [r.height for r in slow.table.rows]
    assert fast.width == slow.width                # 프레임 크기까지 같아야 한다
    assert fast.height == slow.height


def test_table_slide_frame_width_equals_column_sum():
    """열 폭의 합 == 표 폭. set_grid가 프레임 크기를 직접 맞추므로 어긋나면 안 된다."""
    td = _table(n_wafers=8)
    prs, layout = _prs()
    pptgen._table_slide(prs, layout, td)
    frame = next(sh for sh in prs.slides[0].shapes if sh.has_table)
    cols = list(frame.table.columns)
    assert frame.width == sum(c.width for c in cols)
    assert frame.height == sum(r.height for r in frame.table.rows)
    n_lab = len(td.labels())                        # CAT2·CAT3 + item
    assert len({c.width for c in cols[n_lab:]}) == 1  # wafer 열은 폭이 모두 같다


def test_table_slide_does_not_use_column_setter(monkeypatch):
    """열 폭 setter로 되돌아가면 실패한다 — 여기가 wafer 수의 제곱으로 느려진 자리다."""
    from pptx.table import _Column

    calls = []
    orig = _Column.width.fset
    monkeypatch.setattr(
        _Column, "width",
        property(_Column.width.fget,
                 lambda s, v: (calls.append(v), orig(s, v))[1]))
    prs, layout = _prs()
    pptgen._table_slide(prs, layout, _table(n_wafers=8))
    assert calls == []


# ── 2) trend 집계 ────────────────────────────────────────────
def _trend_data() -> dict[str, pl.DataFrame]:
    df = pl.DataFrame({
        "key": [f"k{i}" for i in range(8)],
        "lot": ["L1"] * 8,
        "wafer": ["W1", "W1", "W2", "W2", "W3", "W3", "W4", "W4"],
        "gid": ["g1"] * 4 + ["g2"] * 4,
        "A": [0.3, 0.4, 0.5, 0.6, 0.7, 0.8, 0.9, 1.0],
        "B": [1.2, 1.4, 1.6, 1.8, 2.0, 2.2, 2.4, 2.6],
    })
    return {"g1": df.filter(pl.col("gid") == "g1"),
            "g2": df.filter(pl.col("gid") == "g2")}


def test_trend_aggregates_once_per_group(monkeypatch):
    """wafer_stats 호출은 그룹 수만큼 — item마다 부르면 group_by가 반복된다."""
    seen: list[list[str]] = []
    real = mpl_renderer.wafer_stats

    def spy(df, excluded, aliases, agg="avg"):
        seen.append(list(aliases))
        return real(df, excluded, aliases, agg)

    monkeypatch.setattr(mpl_renderer, "wafer_stats", spy)
    styles = [GroupStyle(gid="g1", name="G1"), GroupStyle(gid="g2", name="G2")]
    spec = PlotSpec(title="T", x="W", y="A,B", type="trend", mode="avg")
    mpl_renderer.render(spec, _trend_data(), styles, _rf(), [], (4.0, 3.0))

    # 점 스트립은 **그룹당 1회**, alias는 한 번에 전부 넘긴다. item 수(2)에
    # 비례해 늘면 예전 경로로 되돌아간 것이다.
    # (라인 쪽 group_representatives는 aggregate 모듈 안에서 직접 부르므로
    #  이 첩자에 잡히지 않는다 — 거기도 이미 alias를 한 번에 넘긴다.)
    assert len(seen) == 2, "그룹 수(2)만큼만 집계해야 한다"
    assert all(a == ["A", "B"] for a in seen), "item마다 집계로 되돌아갔다"


def test_trend_values_unchanged():
    """집계를 한 번에 해도 그려지는 y값은 그대로다."""
    styles = [GroupStyle(gid="g1", name="G1")]
    spec = PlotSpec(title="T", x="W", y="A,B", type="trend", mode="avg")
    fig = mpl_renderer.render(spec, _trend_data(), styles, _rf(), [], (4.0, 3.0))
    ax = fig.axes[0]
    # 점의 **순서**는 고정하지 않는다 — wafer_stats는 polars group_by 결과를
    # 그대로 훑고, 그 순서는 실행마다 달라진다(그림에는 영향이 없다).
    pts = sorted((round(float(x), 6), round(float(y), 6))
                 for c in ax.collections for x, y in c.get_offsets().data)
    # g1: A는 wafer 평균 0.35·0.55(x=W 0.34), B는 1.3·1.7(x=W 0.66)
    assert pts == [(0.34, 0.35), (0.34, 0.55), (0.66, 1.3), (0.66, 1.7)]


# ── 3) xlsx 서식 구간 ────────────────────────────────────────
def test_number_format_boundaries():
    assert number_format(0.999) == "0.000"          # <1
    assert number_format(1.0) == "0.00"             # 1 ≤ v ≤ 10
    assert number_format(10.0) == "0.00"
    assert number_format(10.001) == "0.0"           # >10
    assert number_format(-0.5) == "0.000"           # 부호가 아니라 크기로
    assert number_format(None) is None


def test_format_runs_collapses_same_format():
    runs = format_runs([0.1, 0.2, 0.3, 55.0, 66.0])
    assert runs == [(0, 2, "0.000"), (3, 4, "0.0")]


def test_format_runs_skips_empty_cells():
    """빈 칸은 서식을 걸지 않는다 — 구간에서 빠지되 앞뒤를 잇지도 않는다."""
    assert format_runs([None, None]) == []
    assert format_runs([0.1, None, 0.2]) == [(0, 0, "0.000"), (2, 2, "0.000")]
    assert format_runs([]) == []


def test_flag_runs_only_true_blocks():
    assert flag_runs([False, True, True, False, True]) == [(1, 2), (4, 4)]
    assert flag_runs([False, False]) == []
    assert flag_runs([]) == []


# ── 4) 진행 보고 ─────────────────────────────────────────────
def _report_spec(n_pages: int = 2) -> ReportSpec:
    pages = []
    for i in range(1, n_pages + 1):
        p = PageSpec(number=i, title=f"P{i}")
        p.slots[0] = PlotSpec(title="A-B", x="A", y="B")
        pages.append(p)
    return ReportSpec(report="R", pages=pages)


def test_build_deck_reports_progress():
    """done은 1씩 늘고, 마지막이 total과 같고, **total은 도중에 바뀌지 않는다.**"""
    seen: list[tuple[int, int, str]] = []
    prs = pptgen.build_deck(
        report=_report_spec(2),
        experiments=["", "M1"],
        group_styles_of=lambda exp: [],
        plot_data_of=lambda exp, spec: {},
        tables=[_table(), _table()],
        rf=_rf(),
        log_patterns=[],
        exclusion_log=pl.DataFrame({"key_hash": [], "reason": [],
                                    "created_at": []}),
        meta={"title": "T"},
        on_progress=lambda d, t, label: seen.append((d, t, label)),
    )
    dones = [d for d, _t, _l in seen]
    totals = {t for _d, t, _l in seen}
    assert dones == list(range(1, len(dones) + 1))
    assert len(totals) == 1, "분모가 도중에 바뀌면 막대가 뒤로 간다"
    assert dones[-1] == totals.pop() == len(prs.slides.__iter__.__self__._sldIdLst)
    assert [lb for _d, _t, lb in seen][-1] == "제외 이력"


def test_build_deck_progress_total_counts_split_parts():
    """split 모드는 표가 여러 장으로 갈리는데, 그 장수까지 미리 세어야 한다."""
    seen: list[tuple[int, int, str]] = []
    prs = pptgen.build_deck(
        report=_report_spec(1),
        experiments=[""],
        group_styles_of=lambda exp: [],
        plot_data_of=lambda exp, spec: {},
        # lot 경계로 먼저 끊고(15+15) 그 안에서 12장씩 → 12+3, 12+3 = 4장
        tables=[_table(n_wafers=30)],
        rf=_rf(),
        log_patterns=[],
        exclusion_log=pl.DataFrame({"key_hash": [], "reason": [],
                                    "created_at": []}),
        table_mode="split",
        on_progress=lambda d, t, label: seen.append((d, t, label)),
    )
    assert seen[-1][0] == seen[-1][1] == len(prs.slides.__iter__.__self__._sldIdLst)
    assert sum(1 for _d, _t, lb in seen if lb.startswith("표")) == 4


def test_build_deck_without_progress_is_unchanged():
    """on_progress를 주지 않으면 예전과 똑같이 돈다."""
    kw = {
        "report": _report_spec(1), "experiments": [""],
        "group_styles_of": lambda exp: [],
        "plot_data_of": lambda exp, spec: {},
        "tables": [_table()], "rf": _rf(), "log_patterns": [],
        "exclusion_log": pl.DataFrame({"key_hash": [], "reason": [],
                                       "created_at": []})}
    a = pptgen.build_deck(**kw)
    b = pptgen.build_deck(**kw, on_progress=lambda *_: None)
    assert len(a.slides.__iter__.__self__._sldIdLst) == \
           len(b.slides.__iter__.__self__._sldIdLst)


def test_apply_config_reports_stages(tmp_path):
    """[적용]은 실제로 돌 단계만 분모로 센다 — 빈 설정이면 보고가 없다."""
    from etreport.config.settings import AnalysisConfig
    from etreport.model.session import apply_config
    from etreport.model.state import AppState

    seen: list[tuple[int, int, str]] = []
    rep = apply_config(AppState(), AnalysisConfig(name="x"),
                       on_progress=lambda d, t, lb: seen.append((d, t, lb)))
    assert rep.ok and seen == []

    seen.clear()
    cfg = AnalysisConfig(name="y", db_path=str(tmp_path / "없는.duckdb"))
    apply_config(AppState(), cfg,
                 on_progress=lambda d, t, lb: seen.append((d, t, lb)))
    assert seen == [(0, 1, "DB 읽는 중")]      # 실패한 단계는 완료로 세지 않는다
