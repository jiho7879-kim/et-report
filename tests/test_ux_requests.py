"""사용자 요청 8건 회귀 (2026-08-13).

1 trend의 X축 규격 점선 제거 · 2 리포트 탭에도 그룹 스타일 카드 ·
3 템플릿 예시를 한 파일에 시트로 + 리포메터 WIDTH/LENGTH ·
4 콤보 선택 뒤 팝업이 남지 않게(핸들러 지연) · 5 탐색 X=W/L → trend ·
6 PPT 표지·실험 조건 장표 · 7 plot 해상도·글자 크기 · 8 그룹별 평균
"""
from __future__ import annotations

import os
import time as _time

import polars as pl
import pytest

from tests.conftest import qt_until

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _wait(qapp_cls, loop_cls, ms: int = 400) -> None:
    """콤보 핸들러는 팝업을 닫고 60ms 뒤에 돈다 — 그때까지 루프를 돌린다."""
    end = _time.monotonic() + ms / 1000
    while _time.monotonic() < end:
        qapp_cls.processEvents(loop_cls.AllEvents, 20)



@pytest.fixture(scope="module")
def qapp():
    try:
        from PySide6.QtWidgets import QApplication
    except ImportError:                                # pragma: no cover
        pytest.skip("PySide6 없음")
    app = QApplication.instance() or QApplication([])
    yield app
    app.processEvents()


@pytest.fixture
def demo_state(appdata):
    from etreport import demo
    from etreport.model.state import AppState

    st = AppState()
    demo.load_demo(st)
    return st


def _settle():
    from PySide6.QtCore import QCoreApplication, QEventLoop
    _wait(QCoreApplication, QEventLoop)


# ── 1 trend의 X축 점선 ───────────────────────────────────────
def test_trend_has_no_vertical_spec_lines(demo_state):
    """★ WIDTH·LENGTH 위치의 세로 점선은 그리지 않는다."""
    from etreport.model.specs import PlotSpec
    from etreport.render import mpl_renderer

    st = demo_state
    ys = ",".join(st.aliases()[:2])
    fig = mpl_renderer.render(PlotSpec(type="trend", x="W", y=ys),
                              {"": st.data}, st.groups[:1], st.rf, [], (4, 3))
    ax = fig.axes[0]
    dashed = [ln for ln in ax.lines if ln.get_linestyle() in ("--", "dashed")]
    assert dashed == []


# ── 2 그룹 스타일은 화면마다 갖지 않는다 ─────────────────────
def test_group_style_lives_in_one_shared_section(qapp, demo_state):
    """★ 탐색·리포트가 **같은 섹션 하나**를 본다 — 스타일 규칙이 갈라지지 않게.

    예전에는 두 탭이 각자 `GroupStyleCard`를 하나씩 갖고 있었다(설계 §0 B).
    지금은 워크스페이스가 인스펙터 공용 자리에 하나만 만든다.
    """
    from etreport.config.settings import Settings
    from etreport.model.state import StateBus
    from etreport.ui.analysis_ws import AnalysisWorkspace
    from etreport.ui.widgets.group_section import GroupSection

    ws = AnalysisWorkspace(demo_state, StateBus(), Settings.defaults())
    assert len(ws.findChildren(GroupSection)) == 1
    for tab in ws.tab_widgets():
        assert not hasattr(tab, "style_card")
    ws.deleteLater()


def test_group_style_edit_reaches_the_group(qapp, demo_state):
    from etreport.config.settings import Settings
    from etreport.model.state import StateBus
    from etreport.ui.analysis_ws import AnalysisWorkspace

    ws = AnalysisWorkspace(demo_state, StateBus(), Settings.defaults())
    ws.group_section.select(0)
    ws.group_section.cmb_size.setCurrentText("12")

    assert qt_until(lambda: demo_state.groups[0].size == 12)
    ws.deleteLater()


# ── 3 예시 파일 · WIDTH/LENGTH ───────────────────────────────
def test_reformatter_sample_uses_width_length():
    """★ 리포메터 기하 컬럼은 WIDTH·LENGTH (로더가 읽는 이름과 같아야 한다)."""
    from etreport.data.reformatter import load
    from etreport.export import templates_sample as ts

    cols = ts.reformatter_sample().data.columns
    assert "WIDTH" in cols and "LENGTH" in cols
    assert "W" not in cols and "L" not in cols
    assert load is not None


def test_sample_reformatter_geometry_is_read(fake_sheet):
    from etreport.data.reformatter import load
    from etreport.export import templates_sample as ts

    fake_sheet(ts.reformatter_sample().data)
    rf = load("sample.xlsx")

    real = rf.reals()[0]
    assert real.w == 1.0 and real.l == 0.03


def test_bundle_puts_everything_in_one_file(tmp_path, monkeypatch):
    """★ Excel이 있으면 4종 + 설명이 **한 파일의 시트들**로 저장된다."""
    from etreport.export import templates_sample as ts

    written: dict = {}

    def fake_workbook(samples, path):
        written["path"] = path
        written["sheets"] = [n for s in samples
                             for n in (s.title, f"{s.title} {ts.DESC_SHEET}")]
        path.write_text("x", encoding="utf-8")

    monkeypatch.setattr(ts, "_write_workbook", fake_workbook)
    made = ts.save_all(tmp_path)

    assert made == [tmp_path / "sample_templates.xlsx"]
    assert written["sheets"][:2] == ["리포메터", "리포메터 설명"]
    assert len(written["sheets"]) == 8            # 4종 × (본문 + 설명)


# ── 4 콤보 핸들러 지연 ───────────────────────────────────────
def test_combo_handlers_are_deferred(qapp, demo_state):
    """★ 선택 즉시 무거운 일을 하면 팝업이 안 닫힌다 — 한 박자 미룬다."""
    from etreport.model.state import StateBus
    from etreport.ui.tabs.explore import ExploreTab

    tab = ExploreTab(demo_state, StateBus())
    tab.redraw()

    tab.cmb_scale["y"].setCurrentText("log")
    assert demo_state.explore.logy_mode == "auto"      # 아직 실행 전
    assert qt_until(lambda: demo_state.explore.logy_mode == "log")
    tab.deleteLater()


def test_defer_runs_the_callback(qapp):
    from etreport.ui.tabs.common import defer

    seen = []
    defer(seen.append, 1)
    assert seen == []
    _settle()
    assert seen == [1]


# ── 5 탐색에서 W/L → trend ───────────────────────────────────
@pytest.mark.parametrize("axis", ["W", "L"])
def test_explore_switches_to_trend_for_geometry(qapp, demo_state, axis):
    """★ X에 W·L을 넣으면 trend로 그린다(예전에는 빈 화면)."""
    from etreport.model.state import StateBus
    from etreport.ui.tabs.explore import ExploreTab

    tab = ExploreTab(demo_state, StateBus())
    tab.ed_x.setText(axis)
    tab.ed_y.setText(demo_state.aliases()[0])
    tab._axes_changed()
    tab.redraw()                       # [그리기] — 지연 규약상 버튼으로 그린다

    assert demo_state.explore.type == "trend"
    ax = tab.canvas.figure.axes[0]
    assert ax.get_xlabel() == axis
    assert ax.collections or ax.lines                  # 뭔가 그려졌다
    tab.deleteLater()


def test_explore_goes_back_to_scatter(qapp, demo_state):
    from etreport.model.state import StateBus
    from etreport.ui.tabs.explore import ExploreTab

    tab = ExploreTab(demo_state, StateBus())
    tab.ed_x.setText("W")
    tab._axes_changed()
    tab.ed_x.setText(demo_state.aliases()[0])
    tab._axes_changed()

    assert demo_state.explore.type == "scatter"
    tab.deleteLater()


# ── 6 PPT 표지 · 실험 조건 ───────────────────────────────────
def test_deck_starts_with_a_title_slide(demo_state, tmp_path):
    """★ 표지 → 실험 조건 → plot … 순서."""
    from pptx import Presentation

    from etreport.export import deckbuild

    prs = Presentation(deckbuild.generate(demo_state, str(tmp_path / "d.pptx")))
    texts = [" ".join(sh.text_frame.text for sh in s.shapes
                      if sh.has_text_frame) for s in prs.slides]

    assert "ROOT_LOT_ID" in texts[0] and "PA123" in texts[0]
    assert "실험 조건" in texts[1]
    assert "제외 포인트 이력" in texts[-1]


def test_title_meta_has_the_requested_fields(demo_state):
    from etreport.export.deckbuild import deck_meta

    meta = deck_meta(demo_state)

    assert meta["ROOT_LOT_ID"]
    assert meta["STEP_ID"] or meta["STEP_ID"] == ""     # 데모엔 step이 없을 수 있다
    assert set(meta) >= {"title", "ROOT_LOT_ID", "STEP_ID", "TEMPERATURE"}


def test_split_slide_lists_wafer_conditions(tmp_path):
    from pptx import Presentation

    from etreport.data.reformatter import Reformatter
    from etreport.model.specs import ReportSpec
    from etreport.model.split import SplitMatrix
    from etreport.render import pptgen

    sm = SplitMatrix.from_dataframe(pl.DataFrame({
        "lot": ["PA1", "PA1"], "wafer": ["01", "02"], "M1": ["Base", "Hi"]}))
    prs = pptgen.build_deck(
        report=ReportSpec(report="R"), experiments=[""],
        group_styles_of=lambda e: [], plot_data_of=lambda e, s: {},
        tables=[], rf=Reformatter(), log_patterns=[],
        exclusion_log=pl.DataFrame({"key_hash": [], "reason": [],
                                    "created_at": []}),
        split_rows=sm.wide)
    prs.save(tmp_path / "s.pptx")

    tbl = next(sh.table for sh in Presentation(tmp_path / "s.pptx").slides[0].shapes
               if sh.has_table)
    body = [tbl.cell(r, c).text for r in range(len(tbl.rows))
            for c in range(len(tbl.columns))]
    assert "PA1" in body and "01" in body and "Hi" in body


# ── 7 해상도 · 글자 크기 ─────────────────────────────────────
def test_plot_dpi_and_font_are_bigger_than_before():
    from etreport.render import mpl_renderer, pptgen

    assert pptgen.PLOT_DPI >= 200                      # 예전 150
    assert mpl_renderer.FONT_MIN_PT >= 8               # 예전 6
    small = mpl_renderer._font_size((3.5, 2.6), compact=True)
    big = mpl_renderer._font_size((6.0, 4.0), compact=False)
    assert small >= 8 and big > small


# ── 8 그룹별 평균 ────────────────────────────────────────────
def _grouped_state():
    from etreport.data.reformatter import Reformatter, Rule
    from etreport.model.specs import GroupStyle, ReportSpec, TableRowSpec
    from etreport.model.state import AppState

    st = AppState()
    st.data = pl.DataFrame({
        "key": [str(i) for i in range(4)],
        "lot": ["PA"] * 4, "wafer": ["01", "02", "03", "04"],
        "gid": ["g0", "g0", "g1", "g1"],
        "Vt": [0.40, 0.50, 0.80, 1.00]})
    st.groups = [GroupStyle(gid="g0", name="Base", ref=True),
                 GroupStyle(gid="g1", name="Hi")]
    st.rf = Reformatter(rules=[Rule(category="REAL", itemid="P1", alias="Vt",
                                    absolute=False, scale=1.0, formula="",
                                    unit="V", speclow=None, spechigh=None,
                                    target=None, row=2)])
    st.report = ReportSpec(report="R", table_rows=[TableRowSpec("Vt", ["DC"])])
    return st


def test_group_average_table_has_group_columns():
    """★ 열이 wafer가 아니라 그룹이고, 값은 그룹 평균이다."""
    from etreport.export.excel import SummaryOptions, build_table

    td = build_table(_grouped_state(), "DC", SummaryOptions(agg="gavg"))

    assert td.header_lots == [("그룹", ["Base", "Hi"])]
    assert td.rows[0]["values"] == pytest.approx([0.45, 0.90])


def test_group_average_supports_delta_vs_ref():
    from etreport.export.excel import SummaryOptions, build_table

    td = build_table(_grouped_state(), "DC",
                     SummaryOptions(agg="gavg", delta_vs_ref=True))

    assert td.rows[0]["values"] == pytest.approx([0.0, 0.45])   # REF 대비


def test_summary_tab_offers_three_modes(qapp, demo_state):
    from etreport.model.state import StateBus
    from etreport.ui.tabs.summary import SummaryTab

    tab = SummaryTab(demo_state, StateBus())
    labels = [tab.agg.itemText(i) for i in range(tab.agg.count())]

    assert labels[:3] == ["평균", "Std (wafer 내)", "그룹별 평균"]
    tab.agg.setCurrentIndex(2)
    assert tab.agg_mode() == "gavg" and tab.agg_label() == "그룹별 평균"
    tab.deleteLater()
