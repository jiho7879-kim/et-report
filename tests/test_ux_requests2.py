"""사용자 요청 3건 (2026-08-13, 2차).

1 plot별 점 표시(site/avg/med/std)를 탐색·리포트에서 고른다 ·
2 PPT 뒤쪽에 그룹별 평균 표 페이지 · 3 요약에 '그룹별 wafer' 표
"""
from __future__ import annotations

import os
import time as _time

import polars as pl
import pytest

from tests.conftest import qt_until

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")


def _wait(qapp_cls, loop_cls, ms: int = 400) -> None:
    """콤보 핸들러는 팝업을 닫고 60ms 뒤에 돈다 — 그때까지 루프를 돌린다."""
    end = _time.monotonic() + ms / 1000
    while _time.monotonic() < end:
        qapp_cls.processEvents(loop_cls.AllEvents, 20)



@pytest.fixture(scope="module")
def qapp():
    try:
        from PySide6.QtWidgets import QApplication
    except ImportError:                                # pragma: no cover
        pytest.skip("PySide6 없음")
    app = QApplication.instance() or QApplication([])
    yield app
    app.processEvents()


@pytest.fixture
def demo_state(appdata):
    from etreport import demo
    from etreport.model.state import AppState

    st = AppState()
    demo.load_demo(st)
    return st


def _settle():
    from PySide6.QtCore import QCoreApplication, QEventLoop
    _wait(QCoreApplication, QEventLoop)


# ── 1 점 표시 모드 ───────────────────────────────────────────
def test_point_modes_are_the_template_mode_values():
    """UI 목록이 템플릿 Mode 열과 같은 값이어야 한다(단일 진실)."""
    from etreport.model.specs import POINT_MODES

    assert POINT_MODES == ("site", "avg", "med", "std")


@pytest.mark.parametrize("idx,mode", [(0, "site"), (1, "avg"), (2, "med"),
                                      (3, "std")])
def test_explore_point_combo_sets_mode(qapp, demo_state, idx, mode):
    from etreport.model.state import StateBus
    from etreport.ui.tabs.explore import ExploreTab

    tab = ExploreTab(demo_state, StateBus())
    tab.redraw()
    tab.cmb_point.setCurrentIndex(idx)

    assert qt_until(lambda: demo_state.explore.mode == mode)
    tab.deleteLater()


def test_med_draws_one_point_per_wafer(qapp, demo_state):
    """★ med은 wafer마다 한 점 — 측정점 그대로(site)보다 훨씬 적다."""
    from etreport.model.state import StateBus
    from etreport.render import mpl_renderer
    from etreport.ui.tabs.explore import ExploreTab

    tab = ExploreTab(demo_state, StateBus())
    tab.redraw()
    # 점을 무엇으로 그렸는지는 성능 문제라 렌더러의 `point_xy`로 센다
    site = len(mpl_renderer.point_xy(tab.canvas.figure.axes[0]))

    tab.cmb_point.setCurrentIndex(2)               # wafer 중앙값
    qt_until(lambda: demo_state.explore.mode == "med")
    tab.redraw()                       # [그리기] — 지연 규약상 버튼으로 그린다
    med = len(mpl_renderer.point_xy(tab.canvas.figure.axes[0]))

    n_wafer = demo_state.data.select(["lot", "wafer"]).unique().height
    assert med == n_wafer < site
    tab.deleteLater()


def test_report_inspector_edits_slot_mode(qapp, demo_state):
    """★ 슬롯마다 따로 고를 수 있다(템플릿 Mode 열과 같은 값)."""
    from etreport.model.state import StateBus
    from etreport.ui.tabs.report import ReportTab

    tab = ReportTab(demo_state, StateBus())
    tab.rebuild()
    tab._select(0)
    tab.cmb_point.setCurrentIndex(2)

    assert qt_until(
        lambda: demo_state.report.pages[0].slots[0].mode == "med")
    tab._select(1)                                  # 다른 슬롯은 그대로
    assert tab.cmb_point.currentIndex() == 0
    tab.deleteLater()


def test_slot_mode_survives_reselect(qapp, demo_state):
    from etreport.model.state import StateBus
    from etreport.ui.tabs.report import ReportTab

    tab = ReportTab(demo_state, StateBus())
    tab.rebuild()
    tab._select(0)
    tab.cmb_point.setCurrentIndex(1)
    qt_until(lambda: demo_state.report.pages[0].slots[0].mode == "avg")
    tab._select(1)
    tab._select(0)

    assert tab.cmb_point.currentIndex() == 1        # 고른 값이 그대로 보인다
    tab.deleteLater()


# ── 2 PPT 그룹별 평균 페이지 ─────────────────────────────────
def test_deck_appends_group_average_tables(demo_state, tmp_path):
    """★ 표 전부 → **그룹별 평균 표** → 제외 이력 순."""
    from pptx import Presentation

    from etreport.export import deckbuild

    prs = Presentation(deckbuild.generate(demo_state, str(tmp_path / "d.pptx")))
    titles = [next((sh.text_frame.text for sh in s.shapes
                    if sh.has_text_frame and sh.text_frame.text), "")
              for s in prs.slides]

    cat1 = demo_state.report.table_names()[0]
    # wafer가 많으면 표 제목 뒤에 넘침 안내가 붙는다 — 제목 앞부분으로 본다
    plain = [i for i, t in enumerate(titles)
             if t.startswith(cat1) and "그룹별" not in t]
    grouped = [i for i, t in enumerate(titles) if t.startswith(f"{cat1} — 그룹별")]
    assert plain and grouped and grouped[0] > plain[0]
    assert sum(t.endswith("그룹별 평균") for t in titles) == \
        len(demo_state.report.table_names())
    assert "제외 포인트 이력" in titles[-1]           # 이력은 여전히 마지막


def test_group_tables_are_skipped_without_groups(demo_state):
    from etreport.export.deckbuild import _group_tables

    demo_state.groups = []
    assert _group_tables(demo_state) == []


def test_group_table_columns_are_groups(demo_state):
    from etreport.export.deckbuild import _group_tables

    td = _group_tables(demo_state)[0]
    names = [g.name for g in demo_state.groups]

    assert td.header_lots == [("그룹", names)]
    assert td.name.endswith("그룹별 평균")


# ── 3 그룹별 wafer 표 ────────────────────────────────────────
def _grouped_state():
    from etreport.data.reformatter import Reformatter, Rule
    from etreport.model.specs import GroupStyle, ReportSpec, TableRowSpec
    from etreport.model.state import AppState

    st = AppState()
    st.data = pl.DataFrame({
        "key": [str(i) for i in range(4)],
        "lot": ["PA", "PA", "PB", "PB"],
        "wafer": ["01", "02", "01", "02"],
        "gid": ["g1", "g0", "g0", "g1"],           # 그룹이 lot을 가로지른다
        "Vt": [0.40, 0.50, 0.60, 0.70]})
    st.groups = [GroupStyle(gid="g0", name="Base", ref=True),
                 GroupStyle(gid="g1", name="Hi")]
    st.rf = Reformatter(rules=[Rule(category="REAL", itemid="P1", alias="Vt",
                                    absolute=False, scale=1.0, formula="",
                                    unit="V", speclow=None, spechigh=None,
                                    target=None, row=2)])
    st.report = ReportSpec(report="R", table_rows=[TableRowSpec("Vt", ["DC"])])
    return st


def test_group_wafer_table_groups_the_columns():
    """★ 열은 wafer인데 **그룹 머리글로 묶여** 그룹 순서로 정렬된다."""
    from etreport.export.excel import SummaryOptions, build_table

    td = build_table(_grouped_state(), "DC", SummaryOptions(agg="gwafer"))

    assert td.header_lots == [("Base", ["PA·02", "PB·01"]),
                              ("Hi", ["PA·01", "PB·02"])]
    # Base = (PA,02)=0.50, (PB,01)=0.60 · Hi = (PA,01)=0.40, (PB,02)=0.70
    assert td.rows[0]["values"] == pytest.approx([0.50, 0.60, 0.40, 0.70])


def test_group_wafer_falls_back_to_lots_without_groups():
    from etreport.export.excel import SummaryOptions, build_table

    st = _grouped_state()
    st.groups = []
    st.data = st.data.with_columns(pl.lit("").alias("gid"))

    td = build_table(st, "DC", SummaryOptions(agg="gwafer"))

    assert td.header_lots == [("PA", ["01", "02"]), ("PB", ["01", "02"])]


def test_hidden_group_drops_its_columns():
    from etreport.export.excel import SummaryOptions, build_table

    st = _grouped_state()
    st.groups[1].visible = False

    td = build_table(st, "DC", SummaryOptions(agg="gwafer"))

    assert td.header_lots == [("Base", ["PA·02", "PB·01"])]


def test_summary_tab_offers_four_modes(qapp, demo_state):
    from PySide6.QtWidgets import QTableWidget

    from etreport.model.state import StateBus
    from etreport.ui.tabs.summary import SummaryTab

    tab = SummaryTab(demo_state, StateBus())
    labels = [tab.agg.itemText(i) for i in range(tab.agg.count())]

    assert labels == ["평균", "Std (wafer 내)", "그룹별 평균", "그룹별 wafer"]
    tab.agg.setCurrentIndex(3)
    assert tab.agg_mode() == "gwafer"
    tab.rebuild()                                   # 화면도 그려진다
    assert "그룹별 wafer" in tab.agg_label()
    tbl = next(w for w in tab.findChildren(QTableWidget)
               if w.objectName() == "sumTable")
    heads = [tbl.horizontalHeaderItem(i).text() for i in range(tbl.columnCount())]
    # 라벨 열 수는 CAT 개수(§3.3)와 규격 열(§14)을 따라간다 — 그 뒤가 wafer 열이다
    from etreport.export.excel import SummaryOptions, build_table
    n_label = len(build_table(demo_state, demo_state.report.table_names()[0],
                              SummaryOptions(agg="gwafer")).labels())
    assert all("\n" in h and "·" in h for h in heads[n_label:])  # 그룹\nlot·wafer
    tab.deleteLater()


def test_copy_matches_the_group_wafer_screen(qapp, demo_state):
    """복사(TSV)도 같은 열 구성이어야 한다 — 화면 = 출력."""
    from etreport.export.excel import SummaryOptions, build_table, to_tsv

    opt = SummaryOptions(agg="gwafer")
    td = build_table(_grouped_state(), "DC", opt)
    lines = to_tsv(td, opt).splitlines()

    assert lines[0].split("\t")[-4:] == ["Base", "Base", "Hi", "Hi"]
    assert lines[1].split("\t")[-4:] == ["PA·02", "PB·01", "PA·01", "PB·02"]
